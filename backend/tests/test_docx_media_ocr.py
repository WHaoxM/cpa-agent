"""Contracts for the shared DOCX -> existing OCR sidecar adaptation."""
import io

from docx import Document
from PIL import Image


def _document(tmp_path):
    image = io.BytesIO()
    Image.new("RGB", (400, 100), "white").save(image, "PNG")
    image.seek(0)
    doc = Document()
    doc.add_paragraph("BODY")
    doc.add_picture(image)
    path = tmp_path / "image.docx"
    doc.save(path)
    return str(path)


def test_agent_ingest_and_l0_share_image_transcript(tmp_path, monkeypatch):
    from app.utils import docx_text
    from app.services.agent_runtime.media.ingest import ingest_file
    from app.pipeline.entity_extractor.l0_full_parser import L0FullParser

    calls = []
    def ocr(blob, content_type):
        calls.append(content_type)
        assert blob.startswith(b"\x89PNG")
        return {"success": True, "text": "RASTER_ONLY", "engine": "contract-ocr"}
    monkeypatch.setattr(docx_text, "_ocr_embedded_image", ocr)
    path = _document(tmp_path)
    artifacts, meta, ok = ingest_file(path, "image.docx", purpose="resume_parse")
    full = next(a["body_text"] for a in artifacts if a["kind"] == "full_text")
    assert ok and "BODY" in full and "RASTER_ONLY" in full
    assert meta["docx_media"][0]["status"] == "processed"
    assert meta["accuracy_verified"] is False
    l0 = L0FullParser().parse_file(path)
    assert l0.success and "RASTER_ONLY" in l0.text
    assert len(calls) == 2


def test_disabled_or_failed_ocr_is_partial_and_retryable(tmp_path):
    from app.utils.docx_text import extract_docx_with_media
    path = _document(tmp_path)
    for kwargs in ({"allow_ocr": False}, {"image_ocr": lambda *_: {"success": False}}):
        text, meta = extract_docx_with_media(path, **kwargs)
        assert "BODY" in text
        assert meta["source_coverage"] == "unverified"
        assert meta["retryable"] is True
        assert meta["coverage_issues"]


def test_header_picture_is_not_silently_omitted(tmp_path):
    from app.utils.docx_text import extract_docx_with_media
    path = _document(tmp_path)
    doc = Document(path)
    image = io.BytesIO()
    Image.new("RGB", (100, 100), "red").save(image, "PNG")
    image.seek(0)
    doc.sections[0].header.paragraphs[0].add_run().add_picture(image)
    doc.save(path)
    _, meta = extract_docx_with_media(path, image_ocr=lambda *_: {
        "success": True, "text": "HEADER_OR_BODY", "engine": "contract-ocr"})
    assert len(meta["docx_media"]) == 2
