"""Media coverage is structural, never inferred from resume headings or length."""

from types import SimpleNamespace

from app.services.agent_runtime.media.ingest import extract_text_from_path


def test_pdf_always_verifies_with_l0_regardless_of_apparent_content(tmp_path, monkeypatch):
    pdf = tmp_path / "resume.pdf"
    pdf.write_bytes(b"%PDF-1.4 fixture")
    light = "Education\nSkills\n" + "full-looking body " * 300
    monkeypatch.setattr("app.services.agent_runtime.media.ingest._fallback_pdf",
                        lambda *_: (light, "pypdf"))
    called = []

    class Parser:
        def parse_file(self, path):
            called.append(path)
            return SimpleNamespace(success=True, text="layout verified", blocks=[],
                                   engine="mineru", error=None)

    monkeypatch.setattr("app.pipeline.entity_extractor.l0_full_parser.L0FullParser", Parser)
    text, meta = extract_text_from_path(str(pdf), pdf.name, prefer_light=True, allow_heavy=True)
    assert text == "layout verified"
    assert called == [str(pdf)]
    assert meta["escalated_heavy"] is True
    assert meta["engine"] == "mineru"


def test_pdf_light_fallback_never_claims_complete_source(tmp_path, monkeypatch):
    pdf = tmp_path / "resume.pdf"
    pdf.write_bytes(b"%PDF-1.4 fixture")
    monkeypatch.setattr("app.services.agent_runtime.media.ingest._fallback_pdf",
                        lambda *_: ("Education\nSkills\n" + "body " * 100, "pypdf"))

    class Parser:
        def parse_file(self, _path):
            return SimpleNamespace(success=False, text="", blocks=[], engine="mineru", error="offline")

    monkeypatch.setattr("app.pipeline.entity_extractor.l0_full_parser.L0FullParser", Parser)
    text, meta = extract_text_from_path(str(pdf), pdf.name, prefer_light=True, allow_heavy=True)
    assert text
    assert meta["source_coverage"] == "unverified"
    assert meta["coverage_reason"] == "heavy_parse_failed_light_fallback"


def test_short_docx_without_complex_shapes_is_not_length_classified(tmp_path):
    from docx import Document

    path = tmp_path / "short.docx"
    doc = Document()
    doc.add_paragraph("短文本")
    doc.save(path)
    text, meta = extract_text_from_path(str(path), path.name, prefer_light=True, allow_heavy=True)
    assert text == "短文本"
    assert meta.get("source_coverage") != "unverified"


def test_docx_with_embedded_image_marks_source_unverified(tmp_path, monkeypatch):
    from docx import Document
    from PIL import Image
    from app.services.agent_runtime.media.file_store import _layout_blocks_incomplete
    monkeypatch.setattr("app.utils.docx_text._ocr_embedded_image", lambda *_: {
        "success": False, "error": "OCR unavailable"})

    image = tmp_path / "embedded.png"
    Image.new("RGB", (8, 8), "white").save(image)
    path = tmp_path / "image-text.docx"
    doc = Document()
    doc.add_paragraph("可见正文")
    doc.add_picture(str(image))
    doc.save(path)
    text, meta = extract_text_from_path(str(path), path.name, prefer_light=True, allow_heavy=True)
    assert "可见正文" in text
    assert meta["source_coverage"] == "unverified"
    assert meta["coverage_reason"] == "docx_parts_unprocessed"
    assert _layout_blocks_incomplete({"ingest_meta": meta}) is True
