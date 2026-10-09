"""Agent file / voice media pipeline — not a thin FileReader dump."""

import io
import json
from types import SimpleNamespace

import pytest

from app.services.agent_runtime.fork_jobs import set_inline_for_tests
from app.services.agent_runtime.media.context_mount import (
    build_attachment_mount_card,
    normalize_content_parts,
    user_visible_text,
)
from app.services.agent_runtime.media.file_store import AgentFileService
from app.services.agent_runtime.media.ingest import (
    _fallback_docx,
    _fallback_pdf,
    build_artifacts,
    extract_text_from_path,
    purpose_allows_heavy,
)
from app.services.agent_runtime.media.session_resources import mount_resources, list_resources
from app.services.agent_runtime.media.voice_turn import VoiceTurnService
from app.services.agent_runtime.permissions import visible_tools_for
from app.services.agent_runtime.modes import InteractionMode


@pytest.fixture(autouse=True)
def _inline_ingest():
    """Tests expect ready|failed after upload; run fork handlers inline."""
    set_inline_for_tests(True)
    yield
    set_inline_for_tests(False)


class _MemFile:
    def __init__(self, name: str, data: bytes, mimetype: str = "text/plain"):
        self.filename = name
        self.mimetype = mimetype
        self._data = data

    def save(self, path: str):
        with open(path, "wb") as f:
            f.write(self._data)


def test_chat_purpose_escalates_heavy_by_default(monkeypatch):
    monkeypatch.delenv("CPA_CHAT_INGEST_HEAVY", raising=False)
    assert purpose_allows_heavy("chat_attachment") is True
    assert purpose_allows_heavy("pipeline") is True
    monkeypatch.setenv("CPA_CHAT_INGEST_HEAVY", "0")
    assert purpose_allows_heavy("chat_attachment") is False


def test_full_text_artifact_keeps_content_after_80k():
    text = "正文" * 40_000 + "文末关键结论"
    artifacts = build_artifacts(text, {"engine": "utf8", "ext": "txt"})
    full = next(a for a in artifacts if a["kind"] == "full_text")
    assert full["body_text"].endswith("文末关键结论")
    assert full["meta"]["truncated"] is False


def test_light_pdf_fallback_reads_page_after_40(monkeypatch):
    class FakePage:
        def __init__(self, number):
            self.number = number

        def extract_text(self):
            return f"page-{self.number}"

    class FakeReader:
        pages = [FakePage(number) for number in range(1, 42)]

    monkeypatch.setattr("pypdf.PdfReader", lambda _path: FakeReader())
    text, engine = _fallback_pdf("sample.pdf")
    assert engine == "pypdf"
    assert "page-41" in text


def test_docx_fallback_preserves_body_order_and_header(tmp_path):
    from docx import Document

    document = Document()
    document.sections[0].header.paragraphs[0].text = "页眉联系信息"
    document.add_paragraph("教育经历")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "学校与专业"
    document.add_paragraph("项目经历")
    path = tmp_path / "resume.docx"
    document.save(path)

    text, engine = _fallback_docx(str(path))
    assert engine == "mammoth+python-docx"
    assert "页眉联系信息" in text
    assert text.index("教育经历") < text.index("学校与专业") < text.index("项目经历")


def test_docx_fallback_reads_text_box_in_body_order(tmp_path):
    from docx import Document
    from docx.oxml import OxmlElement

    document = Document()
    document.add_paragraph("教育经历")
    host = document.add_paragraph()
    box = OxmlElement("w:txbxContent")
    paragraph = OxmlElement("w:p")
    run = OxmlElement("w:r")
    text_element = OxmlElement("w:t")
    text_element.text = "文本框中的项目经历"
    run.append(text_element)
    paragraph.append(run)
    box.append(paragraph)
    host.add_run()._r.append(box)
    document.add_paragraph("专业技能")
    path = tmp_path / "boxed-resume.docx"
    document.save(path)

    text, engine = _fallback_docx(str(path))
    assert engine == "mammoth+python-docx"
    assert text.index("教育经历") < text.index("文本框中的项目经历") < text.index("专业技能")

    from app.pipeline.entity_extractor.l0_full_parser import L0FullParser

    pipeline_result = L0FullParser().parse_file(path)
    assert pipeline_result.success is True
    assert pipeline_result.text == text


def test_l0_blank_docx_does_not_claim_parse_success(tmp_path):
    from docx import Document
    from app.pipeline.entity_extractor.l0_full_parser import L0FullParser

    path = tmp_path / "blank.docx"
    Document().save(path)
    result = L0FullParser().parse_file(path)
    assert result.success is False
    assert result.text == ""


def test_chat_extract_skips_l0_when_heavy_disabled(tmp_path, monkeypatch):
    """Opt-out: light-empty must not invoke L0 when allow_heavy=False."""
    pdf = tmp_path / "scan.pdf"
    pdf.write_bytes(b"%PDF-1.4 empty-ish")
    monkeypatch.setattr(
        "app.services.agent_runtime.media.ingest._fallback_pdf",
        lambda *_a, **_k: ("", "pypdf_empty"),
    )

    class BoomParser:
        def __init__(self, *a, **k):
            raise AssertionError("L0/MinerU must not run when allow_heavy=False")

    monkeypatch.setattr(
        "app.pipeline.entity_extractor.l0_full_parser.L0FullParser",
        BoomParser,
        raising=False,
    )
    text, meta = extract_text_from_path(
        str(pdf), "scan.pdf", prefer_light=True, allow_heavy=False,
    )
    assert text == ""
    assert meta.get("engine") == "extract_failed"
    assert meta.get("allow_heavy") is False


def test_chat_extract_escalates_to_l0_when_light_empty(tmp_path, monkeypatch):
    pdf = tmp_path / "scan.pdf"
    pdf.write_bytes(b"%PDF-1.4 empty-ish")
    monkeypatch.setattr(
        "app.services.agent_runtime.media.ingest._fallback_pdf",
        lambda *_a, **_k: ("", "pypdf_empty"),
    )

    class FakeResult:
        success = True
        text = "简历正文来自 L0"
        blocks = []
        engine = "mineru"
        error = None

    class FakeParser:
        def parse_file(self, _path):
            return FakeResult()

    monkeypatch.setattr(
        "app.pipeline.entity_extractor.l0_full_parser.L0FullParser",
        FakeParser,
    )
    text, meta = extract_text_from_path(
        str(pdf), "scan.pdf", prefer_light=True, allow_heavy=True,
    )
    assert "L0" in text
    assert meta.get("escalated_heavy") is True
    assert meta.get("engine") == "mineru"


def test_l0_retains_layout_blocks_after_first_40(tmp_path, monkeypatch):
    pdf = tmp_path / "scan.pdf"
    pdf.write_bytes(b"%PDF-1.4 empty-ish")

    class FakeResult:
        success = True
        text = "简历两页正文"
        blocks = [
            SimpleNamespace(block_id=str(i), block_type="text", text=f"第{i}块", page=i // 40)
            for i in range(68)
        ]
        engine = "mineru"
        error = None

    class FakeParser:
        def parse_file(self, _path):
            return FakeResult()

    monkeypatch.setattr(
        "app.pipeline.entity_extractor.l0_full_parser.L0FullParser",
        FakeParser,
    )
    _, meta = extract_text_from_path(
        str(pdf), "scan.pdf", prefer_light=False, allow_heavy=True,
    )
    assert meta["block_count"] == 68
    assert len(meta["blocks"]) == 68
    assert meta["blocks"][-1]["page"] == 1


def test_mineru_runtime_fault_is_not_sticky():
    from app.pipeline.entity_extractor.pdf_engines import (
        EngineStatus,
        get_mineru_state,
        mark_mineru_failed,
        reset_engine_states,
    )

    reset_engine_states()
    mark_mineru_failed(
        "'PPDocLayoutV2Config' object has no attribute 'reading_order_config'",
        permanent=False,
    )
    st = get_mineru_state()
    assert st.status == EngineStatus.READY
    reset_engine_states()


def test_ingest_reuses_same_sha(tmp_path, monkeypatch):
    set_inline_for_tests(True)
    svc = AgentFileService(mysql_client=SimpleNamespace())
    monkeypatch.setattr(
        "app.services.agent_runtime.media.file_store._upload_root",
        lambda: tmp_path,
    )
    monkeypatch.setattr(
        "app.services.agent_runtime.media.file_store.ensure_media_schema",
        lambda *_a, **_k: False,
    )

    class Boom:
        def execute(self, *a, **k):
            raise RuntimeError("no mysql")

        def fetch_one(self, *a, **k):
            return None

        def fetch_all(self, *a, **k):
            return []

    svc._mysql = Boom()
    body = b"same resume bytes for sha reuse test\n"
    a = svc.upload(
        file_storage=_MemFile("a.txt", body),
        student_id="stu_media",
    )
    assert a["status"] == "ready"
    b = svc.upload(
        file_storage=_MemFile("b.txt", body),
        student_id="stu_media",
    )
    assert b["status"] == "ready"
    assert b["file_id"] != a["file_id"]
    meta_b = json.loads((tmp_path / f"{b['file_id']}.meta.json").read_text(encoding="utf-8"))
    assert meta_b.get("ingest_meta", {}).get("reused_from") == a["file_id"]


def test_ingest_does_not_reuse_sha_across_owners(tmp_path, monkeypatch):
    set_inline_for_tests(True)
    svc = AgentFileService(mysql_client=SimpleNamespace())
    monkeypatch.setattr(
        "app.services.agent_runtime.media.file_store._upload_root",
        lambda: tmp_path,
    )
    monkeypatch.setattr(
        "app.services.agent_runtime.media.file_store.ensure_media_schema",
        lambda *_a, **_k: False,
    )

    class Boom:
        def execute(self, *a, **k):
            raise RuntimeError("no mysql")

        def fetch_one(self, *a, **k):
            return None

        def fetch_all(self, *a, **k):
            return []

    svc._mysql = Boom()
    body = b"same bytes but different owner\n"
    a = svc.upload(file_storage=_MemFile("a.txt", body), student_id="stu_a")
    b = svc.upload(file_storage=_MemFile("b.txt", body), student_id="stu_b")

    meta_b = json.loads((tmp_path / f"{b['file_id']}.meta.json").read_text(encoding="utf-8"))
    assert a["status"] == "ready"
    assert b["status"] == "ready"
    assert meta_b.get("ingest_meta", {}).get("reused_from") is None


def test_mount_resources_requires_student_scope():
    with pytest.raises(PermissionError):
        mount_resources("sess_missing_scope", [{"file_id": "file_x"}], student_id=None)


def test_pdf_upload_returns_before_heavy_work(tmp_path, monkeypatch):
    """HTTP path must not wait on ingest when not inline — status stays ingesting."""
    set_inline_for_tests(False)
    svc = AgentFileService(mysql_client=SimpleNamespace())
    monkeypatch.setattr(
        "app.services.agent_runtime.media.file_store._upload_root",
        lambda: tmp_path,
    )
    monkeypatch.setattr(
        "app.services.agent_runtime.media.file_store.ensure_media_schema",
        lambda *_a, **_k: False,
    )

    class Boom:
        def execute(self, *a, **k):
            raise RuntimeError("no mysql")

        def fetch_one(self, *a, **k):
            return None

        def fetch_all(self, *a, **k):
            return []

    svc._mysql = Boom()
    monkeypatch.setattr(
        svc,
        "complete_ingest",
        lambda *_a, **_k: (_ for _ in ()).throw(AssertionError("sync ingest")),
    )
    enq = []

    def _enq(session_id, kind, payload=None):
        enq.append((kind, payload))
        return kind

    monkeypatch.setattr(
        "app.services.agent_runtime.fork_jobs.enqueue",
        _enq,
    )

    meta = svc.upload(
        file_storage=_MemFile("resume.pdf", b"%PDF-1.4 fake", "application/pdf"),
        student_id="stu_media",
        session_id="sess_media",
        purpose="chat_attachment",
    )
    assert meta["status"] == "ingesting"
    assert enq and enq[0][0] == "media_ingest"
    set_inline_for_tests(True)


def test_ingest_artifacts_redact_and_chunk():
    arts = build_artifacts(
        "name=张三 phone=13800138000 email=a@b.com\n" + ("技能Python " * 200),
        {"engine": "utf8", "blocks": []},
    )
    kinds = {a["kind"] for a in arts}
    assert "excerpt" in kinds and "full_text" in kinds and "chunk" in kinds
    full = next(a for a in arts if a["kind"] == "full_text")["body_text"]
    assert "13800138000" not in full
    assert "a@b.com" not in full


def test_upload_mount_card_no_full_body(tmp_path, monkeypatch):
    svc = AgentFileService(mysql_client=SimpleNamespace())
    monkeypatch.setattr(
        "app.services.agent_runtime.media.file_store._upload_root",
        lambda: tmp_path,
    )
    monkeypatch.setattr(
        "app.services.agent_runtime.media.file_store.ensure_media_schema",
        lambda *_a, **_k: False,
    )

    class Boom:
        def execute(self, *a, **k):
            raise RuntimeError("no mysql")

        def fetch_one(self, *a, **k):
            return None

        def fetch_all(self, *a, **k):
            return []

    svc._mysql = Boom()
    meta = svc.upload(
        file_storage=_MemFile("resume.txt", b"hello CPA student phone=13900139000"),
        student_id="stu_media",
        session_id="sess_media_unique",
    )
    assert meta["status"] == "ready"
    assert meta["file_id"].startswith("file_")
    assert "13900139000" not in (meta.get("excerpt") or "")
    # disk catalog must exist (Claude durable file_id)
    assert (tmp_path / f"{meta['file_id']}.meta.json").is_file()

    mount_resources(
        "sess_media_unique",
        [{"file_id": meta["file_id"]}],
        student_id="stu_media",
    )
    card = build_attachment_mount_card(
        "sess_media_unique",
        turn_file_ids=[meta["file_id"]],
    )
    assert meta["file_id"] in card
    assert "attachment_read" in card
    body = svc.read_for_model(meta["file_id"], max_chars=500)
    assert "hello CPA" in body
    assert "13900139000" not in body


def test_corrupt_docx_marks_failed_not_ready(tmp_path, monkeypatch):
    svc = AgentFileService(mysql_client=SimpleNamespace())
    monkeypatch.setattr(
        "app.services.agent_runtime.media.file_store._upload_root",
        lambda: tmp_path,
    )
    monkeypatch.setattr(
        "app.services.agent_runtime.media.file_store.ensure_media_schema",
        lambda *_a, **_k: False,
    )

    class Boom:
        def execute(self, *a, **k):
            raise RuntimeError("no mysql")

        def fetch_one(self, *a, **k):
            return None

        def fetch_all(self, *a, **k):
            return []

    svc._mysql = Boom()
    meta = svc.upload(
        file_storage=_MemFile("bad.docx", b"PK fake not a docx"),
        student_id="stu_media",
    )
    assert meta["file_id"].startswith("file_")
    assert meta["status"] == "failed"
    assert meta.get("error")


def test_content_parts_user_visible_not_dump():
    parts = normalize_content_parts(
        "请分析这份简历",
        [{"type": "document", "file_id": "file_abc", "title": "resume.pdf"}],
        ["file_abc"],
    )
    visible = user_visible_text(parts)
    assert "请分析这份简历" in visible
    assert "[附件:" in visible
    assert "PDF正文" not in visible


def test_voice_turn_lifecycle():
    vt = VoiceTurnService(mysql_client=SimpleNamespace())

    class Boom:
        def execute(self, *a, **k):
            raise RuntimeError("no mysql")

        def fetch_one(self, *a, **k):
            return None

    vt._mysql = Boom()
    turn = vt.start(session_id="sess_v", student_id="stu_v", lang="zh-CN")
    assert turn["status"] == "capturing"
    mid = vt.update(turn["turn_id"], interim_transcript="我想做")
    assert mid["status"] == "transcribing"
    ready = vt.update(turn["turn_id"], transcript="我想做数据分析", is_final=True)
    assert ready["status"] == "ready"
    part = vt.as_content_part(turn["turn_id"])
    assert part and part["type"] == "voice_transcript"
    assert "数据分析" in part["text"]


def test_ask_attachment_read_only_when_allowlisted():
    assert "attachment_read" not in visible_tools_for(InteractionMode.ASK, "qa")
    assert visible_tools_for(
        InteractionMode.ASK, "qa", plan_tool_allowlist={"attachment_read"},
    ) == {"attachment_read"}
    assert "attachment_read" in visible_tools_for(InteractionMode.AGENT, "qa")


def test_prepare_user_turn_stm_markers(tmp_path, monkeypatch):
    from app.services.agent_runtime.media.prepare_turn import prepare_user_turn
    from app.services.agent_runtime.media.file_store import AgentFileService

    monkeypatch.setattr(
        "app.services.agent_runtime.media.file_store._upload_root",
        lambda: tmp_path,
    )
    monkeypatch.setattr(
        "app.services.agent_runtime.media.file_store.ensure_media_schema",
        lambda *_a, **_k: False,
    )
    monkeypatch.setattr(
        "app.services.agent_runtime.media.session_resources._mount_root",
        lambda: tmp_path / "mounts",
    )

    class Boom:
        def execute(self, *a, **k):
            raise RuntimeError("no mysql")

        def fetch_one(self, *a, **k):
            return None

        def fetch_all(self, *a, **k):
            return []

    svc = AgentFileService(mysql_client=Boom())
    meta = svc.upload(
        file_storage=_MemFile("cv.txt", b"resume body for CPA"),
        student_id="stu_p",
        session_id="sess_p",
    )
    out = prepare_user_turn(
        "请分析",
        session_id="sess_p",
        student_id="stu_p",
        attachment_ids=[meta["file_id"]],
        mount=True,
    )
    assert "请分析" in out["visible_text"]
    assert "[附件:" in out["visible_text"]
    assert meta["file_id"] in out["turn_file_ids"]
    assert out["attach_card"]
